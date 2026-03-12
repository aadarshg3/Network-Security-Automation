import logging
import os
from openpyxl import load_workbook

logger = logging.getLogger("sdwan.excel")


def _sheet(wb, name, data_start_row=2):
    ws = wb[name]
    data = {}
    for key, value in ws.iter_rows(min_row=data_start_row, max_col=2, values_only=True):
        if key is None or str(key).strip() == "":
            continue
        # Skip any column-header rows (Name/Value, Variable Name/Value)
        if str(key).strip().lower() in ("name", "variable name"):
            continue
        data[str(key).strip()] = str(value).strip() if value is not None else ""
    return data


def load(path="configs/newvalues.xlsx"):
    if not os.path.isfile(path):
        raise FileNotFoundError(f"Excel not found: {path}")

    wb = load_workbook(filename=path, data_only=True)
    required = ["connection", "uc1_ids", "uc1_vars", "uc2_ids", "uc2_vars", "uc3_ids", "uc3_vars"]
    missing = [s for s in required if s not in wb.sheetnames]
    if missing:
        raise ValueError(f"Excel missing sheets: {missing}")

    conn     = _sheet(wb, "connection")
    uc1_ids  = _sheet(wb, "uc1_ids")
    uc1_vars = _sheet(wb, "uc1_vars")
    uc2_ids  = _sheet(wb, "uc2_ids")
    uc2_vars = _sheet(wb, "uc2_vars")
    uc3_ids  = _sheet(wb, "uc3_ids")
    uc3_vars = _sheet(wb, "uc3_vars")

    if not conn.get("vmanage_host"):
        raise ValueError("connection.vmanage_host is empty in Excel")

    return {
        "vmanage": {
            "host":     conn["vmanage_host"],
            "username": conn.get("username", "admin"),
        },
        "uc1": {
            "device_template_id":  uc1_ids.get("DEVICE_TEMPLATE_ID", ""),
            "device_uuid":         uc1_ids.get("DEVICE_UUID", ""),
            "tracker_template_id": uc1_ids.get("TRACKER_TEMPLATE_ID", ""),
            "nat_template_id":     uc1_ids.get("NAT_TEMPLATE_ID", ""),
            "variables":           uc1_vars,
        },
        "uc2": {
            "device_template_id": uc2_ids.get("DEVICE_TEMPLATE_ID", ""),
            "device_uuid":        uc2_ids.get("DEVICE_UUID", ""),
            "sig_template_id":    uc2_ids.get("SIG_TEMPLATE_ID", ""),
            "sig_cred_id":        uc2_ids.get("CISCO_SIG_CRED", ""),
            "variables":          uc2_vars,
        },
        "uc3": {
            "device_template_id": uc3_ids.get("DEVICE_TEMPLATE_ID", ""),
            "device_uuid":        uc3_ids.get("DEVICE_UUID", ""),
            "pepguest_vpn_id":    uc3_ids.get("PEPGUEST_VPN_TEMPLATE_ID", ""),
            "pepguest_subif_id":  uc3_ids.get("PEPGUEST_SUBIF_TEMPLATE_ID", ""),
            "pepinet_vpn_id":     uc3_ids.get("PEPINET_VPN_TEMPLATE_ID", ""),
            "pepinet_subif_id":   uc3_ids.get("PEPINET_SUBIF_TEMPLATE_ID", ""),
            "corporate_vpn_id":   uc3_ids.get("CORPORATE_VPN_TEMPLATE_ID", ""),
            "corporate_subif_id": uc3_ids.get("CORPORATE_SUBIF_TEMPLATE_ID", ""),
            "corporate_bgp_id":   uc3_ids.get("CORPORATE_BGP_TEMPLATE_ID", ""),
            "variables":          uc3_vars,
        },
    }
