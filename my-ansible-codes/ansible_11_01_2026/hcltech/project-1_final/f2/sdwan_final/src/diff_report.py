import difflib
import logging
import os
from datetime import datetime

logger = logging.getLogger("sdwan.diff")


def show_and_save(before, after, output_dir="outputs"):
    os.makedirs(output_dir, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")

    diff_lines = list(difflib.unified_diff(
        before.splitlines(keepends=True),
        after.splitlines(keepends=True),
        fromfile="BEFORE",
        tofile="AFTER"
    ))
    diff_text = "".join(diff_lines) if diff_lines else "No differences detected\n"

    print("\nCONFIG DIFF")
    print("-" * 60)
    shown = diff_text.splitlines()[:80]
    for line in shown:
        print(line)
    remaining = len(diff_text.splitlines()) - len(shown)
    if remaining > 0:
        print(f"  ... {remaining} more lines in diff file")

    paths = {
        "pre":  os.path.join(output_dir, f"{ts}_pre.txt"),
        "post": os.path.join(output_dir, f"{ts}_post.txt"),
        "diff": os.path.join(output_dir, f"{ts}_diff.txt"),
    }
    with open(paths["pre"],  "w") as f: f.write(before)
    with open(paths["post"], "w") as f: f.write(after)
    with open(paths["diff"], "w") as f: f.write(diff_text)

    excel_path = _save_excel(before, after, output_dir, ts)
    if excel_path:
        paths["excel"] = excel_path

    logger.info(f"Diff saved prefix={ts}")
    return paths


def _save_excel(before, after, output_dir, ts):
    try:
        import pandas as pd
        from openpyxl.styles import PatternFill
    except ImportError:
        return ""

    path   = os.path.join(output_dir, f"{ts}_diff_report.xlsx")
    diffs  = list(difflib.Differ().compare(before.splitlines(keepends=True), after.splitlines(keepends=True)))
    data   = [{"Change": d[0], "Line": d[2:].strip("\n")} for d in diffs]
    df     = pd.DataFrame(data)
    writer = pd.ExcelWriter(path, engine="openpyxl")
    df.to_excel(writer, sheet_name="Diff", index=False)
    ws    = writer.sheets["Diff"]
    red   = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    green = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    for i, val in enumerate(df["Change"], 2):
        if val == "-":
            for cell in ws[i]: cell.fill = red
        elif val == "+":
            for cell in ws[i]: cell.fill = green
    writer.close()
    logger.info(f"Excel diff saved: {path}")
    return path
