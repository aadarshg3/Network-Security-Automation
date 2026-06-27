"""
src/services/report_service.py

Excel report generation service.
Generates a workbook with:
  - Summary sheet (all UCs, status, timestamps)
  - Per-UC sheets (before/after config, diff, variables, warnings)

The customer's EXACT use case names are preserved in all report content.
"""

from __future__ import annotations

import logging
import os
from datetime import datetime
from typing import List

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment, Border, Font, PatternFill, Side
)
from openpyxl.utils import get_column_letter

from src.models.result_models import CutoverReport, UCResult, UCStatus

logger = logging.getLogger("sdwan.report")

# ── Colour palette ──────────────────────────────────────────────────────
_DARK_BLUE  = "1F4E79"
_MED_BLUE   = "2E75B6"
_LIGHT_BLUE = "DEEAF1"
_GREEN      = "E2EFDA"
_RED        = "FFDCE1"
_YELLOW     = "FFF2CC"
_ORANGE     = "FCE4D6"
_GREY       = "F2F2F2"
_WHITE      = "FFFFFF"

_STATUS_COLOURS = {
    UCStatus.SUCCESS:      _GREEN,
    UCStatus.WARNING:      _YELLOW,
    UCStatus.FAILED:       _RED,
    UCStatus.ROLLED_BACK:  _ORANGE,
    UCStatus.SKIPPED:      _GREY,
    UCStatus.DRY_RUN:      _LIGHT_BLUE,
}


class ReportService:
    """Generates Excel reports for a full cutover run."""

    def __init__(self, reports_dir: str = "outputs"):
        self.reports_dir = reports_dir
        os.makedirs(reports_dir, exist_ok=True)

    def generate(self, report: CutoverReport) -> str:
        """
        Generate and save the Excel report.
        Returns the output file path.
        """
        ts       = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
        filename = f"sdwan_report_{report.site_name}_{ts}.xlsx"
        path     = os.path.join(self.reports_dir, filename)

        wb = Workbook()
        self._build_summary_sheet(wb, report)
        for result in report.results:
            self._build_uc_sheet(wb, result)

        # Remove the default empty sheet
        if "Sheet" in wb.sheetnames:
            del wb["Sheet"]

        wb.save(path)
        logger.info(f"Report saved: {path}")
        return path

    # ──────────────────────────────────────────────────────────────────────
    # Summary Sheet
    # ──────────────────────────────────────────────────────────────────────

    def _build_summary_sheet(self, wb: Workbook, report: CutoverReport) -> None:
        ws = wb.create_sheet("Summary", 0)

        # Title
        ws.merge_cells("A1:G1")
        _cell(ws, "A1", "SD-WAN Automation Report — Site Cutover",
              bold=True, size=14, fg=_DARK_BLUE, colour=_WHITE)
        ws.row_dimensions[1].height = 24

        # Metadata
        meta = [
            ("Site Name",       report.site_name),
            ("Run At (UTC)",    report.run_at.strftime("%Y-%m-%d %H:%M:%S")),
            ("Mode",            "DRY-RUN" if report.dry_run else "LIVE DEPLOY"),
            ("Force",           str(report.force)),
            ("Profile",         report.profile_id or "default"),
            ("Region",          report.region),
            ("Metallic Type",   report.metallic_type),
            ("Design Type",     report.design_type or "N/A"),
            ("Overall Status",  report.overall_status),
        ]
        for r, (key, val) in enumerate(meta, start=2):
            _cell(ws, f"A{r}", key, bold=True, fg=_LIGHT_BLUE)
            _cell(ws, f"B{r}", val)
        ws.column_dimensions["A"].width = 20
        ws.column_dimensions["B"].width = 30

        # UC Results table
        hdr_row = len(meta) + 3
        headers = [
            "Use Case", "Status", "Template ID", "Action ID",
            "Deployed At", "Rollback", "Warnings"
        ]
        for c, h in enumerate(headers, start=1):
            _cell(ws, f"{get_column_letter(c)}{hdr_row}", h,
                  bold=True, fg=_MED_BLUE, colour=_WHITE)
        ws.row_dimensions[hdr_row].height = 18

        for r, result in enumerate(report.results, start=hdr_row + 1):
            colour = _STATUS_COLOURS.get(result.status, _WHITE)
            row_data = [
                result.uc_name,           # exact customer name
                result.display_status,
                result.template_id[:16] + "..." if result.template_id else "",
                result.action_id[:16] + "..." if result.action_id else "",
                result.timestamp.strftime("%Y-%m-%d %H:%M:%S"),
                result.rollback_status or "N/A",
                "; ".join(result.warnings) if result.warnings else "",
            ]
            for c, val in enumerate(row_data, start=1):
                cell = ws[f"{get_column_letter(c)}{r}"]
                cell.value     = val
                cell.fill      = PatternFill("solid", fgColor=colour)
                cell.font      = Font(size=10)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.border    = _thin_border()
            ws.row_dimensions[r].height = 20

        col_widths = [55, 12, 20, 20, 22, 15, 50]
        for c, w in enumerate(col_widths, start=1):
            ws.column_dimensions[get_column_letter(c)].width = w

    # ──────────────────────────────────────────────────────────────────────
    # Per-UC Sheet
    # ──────────────────────────────────────────────────────────────────────

    def _build_uc_sheet(self, wb: Workbook, result: UCResult) -> None:
        # Sheet name must be ≤ 31 chars
        short_name = result.uc_key.upper()
        ws = wb.create_sheet(short_name)

        # Title — exact customer name
        ws.merge_cells("A1:C1")
        _cell(ws, "A1", result.uc_name,
              bold=True, size=12, fg=_DARK_BLUE, colour=_WHITE)
        ws.row_dimensions[1].height = 22

        r = 2
        # Status block
        colour = _STATUS_COLOURS.get(result.status, _WHITE)
        for label, value in [
            ("Site",         result.site_name),
            ("Status",       result.display_status),
            ("Template ID",  result.template_id),
            ("Action ID",    result.action_id or "N/A"),
            ("Timestamp",    result.timestamp.strftime("%Y-%m-%d %H:%M:%S")),
            ("Rollback",     result.rollback_status or "N/A"),
            ("Dry-run",      str(result.dry_run)),
            ("Force",        str(result.force)),
        ]:
            _cell(ws, f"A{r}", label, bold=True, fg=_LIGHT_BLUE)
            cell = ws[f"B{r}"]
            cell.value = value
            cell.fill  = PatternFill("solid", fgColor=colour)
            r += 1

        # Variables
        r += 1
        _cell(ws, f"A{r}", "Variables Applied", bold=True, fg=_MED_BLUE, colour=_WHITE)
        ws.merge_cells(f"A{r}:C{r}")
        r += 1
        for k, v in result.variables_applied.items():
            _cell(ws, f"A{r}", str(k), bold=True)
            _cell(ws, f"B{r}", str(v))
            r += 1

        # Warnings
        if result.warnings:
            r += 1
            _cell(ws, f"A{r}", "Warnings", bold=True, fg=_YELLOW)
            r += 1
            for w in result.warnings:
                _cell(ws, f"A{r}", w, fg=_YELLOW)
                ws.merge_cells(f"A{r}:C{r}")
                r += 1

        # Errors
        if result.errors:
            r += 1
            _cell(ws, f"A{r}", "Errors", bold=True, fg=_RED)
            r += 1
            for e in result.errors:
                _cell(ws, f"A{r}", e, fg=_RED)
                ws.merge_cells(f"A{r}:C{r}")
                r += 1

        # Before config
        r += 1
        _cell(ws, f"A{r}", "Before Config", bold=True, fg=_LIGHT_BLUE)
        r += 1
        _multiline(ws, r, "A", result.before_config or "(not captured)")
        r += _count_lines(result.before_config) + 2

        # After config
        _cell(ws, f"A{r}", "After Config", bold=True, fg=_LIGHT_BLUE)
        r += 1
        _multiline(ws, r, "A", result.after_config or "(not captured)")
        r += _count_lines(result.after_config) + 2

        # Diff summary
        _cell(ws, f"A{r}", "Diff Summary", bold=True, fg=_LIGHT_BLUE)
        r += 1
        _multiline(ws, r, "A", result.diff_summary or "(no diff captured)")

        ws.column_dimensions["A"].width = 60
        ws.column_dimensions["B"].width = 60
        ws.column_dimensions["C"].width = 30


# ── Cell helpers ───────────────────────────────────────────────────────────

def _cell(
    ws,
    addr: str,
    value,
    bold: bool = False,
    size: int = 10,
    fg: str = _WHITE,
    colour: str = "000000",
) -> None:
    c = ws[addr]
    c.value     = value
    c.font      = Font(bold=bold, size=size, color=colour)
    c.fill      = PatternFill("solid", fgColor=fg)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    c.border    = _thin_border()


def _thin_border() -> Border:
    thin = Side(style="thin", color="AAAAAA")
    return Border(left=thin, right=thin, top=thin, bottom=thin)


def _multiline(ws, row: int, col: str, text: str) -> None:
    c = ws[f"{col}{row}"]
    c.value     = text
    c.font      = Font(size=9, name="Courier New")
    c.alignment = Alignment(wrap_text=True, vertical="top")
    line_count  = _count_lines(text)
    ws.row_dimensions[row].height = max(15, min(line_count * 13, 200))


def _count_lines(text: str) -> int:
    if not text:
        return 1
    return max(1, len(text.splitlines()))
