"""
Excel-driven configuration loader for the customer package.

The workbook format follows the existing Engineer-1 package so customers can
keep supplying inputs in Excel instead of YAML.
"""

from __future__ import annotations

import os
from dataclasses import dataclass
from typing import Dict, Iterable, List, Tuple

import openpyxl

from src.models.site_models import SiteConfig, SiteProfile


DEFAULT_EXCEL_PATH = os.path.join("configs", "sdwan_config.xlsx")
HEADER_ROWS = {"Field", "Variable Key", "Variable Key (do not change)", "Setting"}


@dataclass
class ExcelConfigBundle:
    settings: Dict
    site: SiteConfig
    profile: SiteProfile
    warnings: List[str]
    excel_path: str


def load_excel_config(excel_path: str = DEFAULT_EXCEL_PATH) -> ExcelConfigBundle:
    path = os.path.abspath(excel_path)
    if not os.path.isfile(path):
        raise FileNotFoundError(f"Excel config not found: {path}")

    wb = openpyxl.load_workbook(path, data_only=True)
    warnings: List[str] = []

    connection = dict(_iter_kv(_find_sheet(wb, "connection")))
    u1_ids = dict(_iter_kv(_find_sheet(wb, "UC1 IDs")))
    u1_vars = dict(_iter_kv(_find_sheet(wb, "UC1 Variables")))
    u2_ids = dict(_iter_kv(_find_sheet(wb, "UC2 IDs")))
    u2_vars = dict(_iter_kv(_find_sheet(wb, "UC2 Variables")))
    u3_ids = dict(_iter_kv(_find_sheet(wb, "UC3 IDs")))
    u3_vars = dict(_iter_kv(_find_sheet(wb, "UC3 Variables")))
    pepguest_ids, pepinet_ids, pepguest_vars, pepinet_vars = _parse_uc4(_find_sheet(wb, "UC4 Variables"))

    vmanage_host = _clean(connection.get("vmanage_host"))
    username = _clean(connection.get("username")) or "admin"
    if not vmanage_host:
        warnings.append("connection.vmanage_host")

    device_template_id = _first_non_empty(
        _clean(u1_ids.get("DEVICE_TEMPLATE_ID")),
        _clean(u2_ids.get("DEVICE_TEMPLATE_ID")),
        _clean(u3_ids.get("DEVICE_TEMPLATE_ID")),
    )
    device_uuid = _first_non_empty(
        _clean(u1_ids.get("DEVICE_UUID")),
        _clean(u2_ids.get("DEVICE_UUID")),
        _clean(u3_ids.get("DEVICE_UUID")),
    )
    if not device_template_id:
        warnings.append("common.device_template_id")
    if not device_uuid:
        warnings.append("common.device_uuid")

    tracker_src_ip = _first_non_empty(
        _clean(u1_vars.get("/0/0/vpn-instance/tracker-src-ip")),
        _clean(u1_vars.get("tracker_src_ip")),
    )

    fw_physical_port = _first_non_empty(
        _clean(u3_ids.get("FW_PHYSICAL_PORT")),
        _clean(u3_vars.get("/0/VPN0_INT_Zscaler_interface_NAME/interface/if-name")),
    )
    fw_transit_vlan = _to_int(_clean(u3_ids.get("FW_TRANSIT_VLAN")), default=50)

    site = SiteConfig.from_dict(
        {
            "site_name": _first_non_empty(_clean(connection.get("site_name")), "excel-site"),
            "device_name": _first_non_empty(_clean(connection.get("device_name")), "cedge"),
            "system_ip": _first_non_empty(_clean(connection.get("system_ip")), ""),
            "device_uuid": device_uuid,
            "region": _first_non_empty(_clean(connection.get("region")), "DEFAULT"),
            "metallic_type": _first_non_empty(_clean(connection.get("metallic_type")), "DEFAULT"),
            "design_type": _clean(connection.get("design_type")) or None,
            "device_template_id": device_template_id,
            "uc1_device_uuid": _first_non_empty(_clean(u1_ids.get("DEVICE_UUID")), device_uuid),
            "uc1_device_template_id": _first_non_empty(_clean(u1_ids.get("DEVICE_TEMPLATE_ID")), device_template_id),
            "uc2_device_uuid": _first_non_empty(_clean(u2_ids.get("DEVICE_UUID")), device_uuid),
            "uc2_device_template_id": _first_non_empty(_clean(u2_ids.get("DEVICE_TEMPLATE_ID")), device_template_id),
            "uc3_device_uuid": _first_non_empty(_clean(u3_ids.get("DEVICE_UUID")), device_uuid),
            "uc3_device_template_id": _first_non_empty(_clean(u3_ids.get("DEVICE_TEMPLATE_ID")), device_template_id),
            "uc4_device_uuid": _first_non_empty(_clean(u3_ids.get("DEVICE_UUID")), device_uuid),
            "uc4_device_template_id": _first_non_empty(_clean(u3_ids.get("DEVICE_TEMPLATE_ID")), device_template_id),
            "tracker_template_id": _clean(u1_ids.get("TRACKER_TEMPLATE_ID")),
            "nat_template_id": _clean(u1_ids.get("NAT_TEMPLATE_ID")),
            "tracker_src_ip": tracker_src_ip,
            "sig_template_id": _clean(u2_ids.get("SIG_TEMPLATE_ID")),
            "cisco_sig_cred_template_id": _clean(u2_ids.get("CISCO_SIG_CRED_ID")),
            "sig_variables": _drop_empty(u2_vars),
            "fw_physical_port": fw_physical_port,
            "fw_transit_vlan": fw_transit_vlan,
            "vpn0_noip_template_id": _clean(u3_ids.get("VPN0_NOIP_TEMPLATE_ID")),
            "firewall_transit_template_id": _clean(u3_ids.get("FW_TRANSIT_SUBIF_TEMPLATE_ID")),
            "uc3_variables": _drop_empty(
                {
                    **_inject_uc3_logical_port(u3_vars, fw_physical_port, fw_transit_vlan),
                    **pepguest_vars,
                    **pepinet_vars,
                }
            ),
            "pepguest_vpn_template_id": _clean(pepguest_ids.get("PEPGUEST_VPN_TEMPLATE_ID")),
            "pepguest_subif_template_id": _clean(pepguest_ids.get("PEPGUEST_SUBIF_TEMPLATE_ID")),
            "pepinet_vpn_template_id": _clean(pepinet_ids.get("PEPINET_VPN_TEMPLATE_ID")),
            "pepinet_subif_template_id": _clean(pepinet_ids.get("PEPINET_SUBIF_TEMPLATE_ID")),
            "uc4_variables": _drop_empty(
                {
                    **_inject_uc3_logical_port(u3_vars, fw_physical_port, fw_transit_vlan),
                    **pepguest_vars,
                    **pepinet_vars,
                }
            ),
        }
    )

    profile = SiteProfile(
        profile_id="excel-default",
        region=site.region,
        metallic_type=site.metallic_type,
        design_type=site.design_type,
        description="Excel-driven customer profile",
        enabled_use_cases=["uc1", "uc2", "uc3", "uc4"],
    )

    settings = {
        "vmanage": {
            "host": vmanage_host,
            "username": username,
            "verify_ssl": False,
            "timeout_seconds": 45,
            "retry_attempts": 4,
            "retry_delay_seconds": 5.0,
            "poll_timeout_seconds": 180,
            "poll_interval_seconds": 5,
        },
        "paths": {
            "state_file": "sdwan_state.json",
            "backups_dir": "backups",
            "logs_dir": "logs",
        },
        "logging": {
            "level": "INFO",
            "format": "%(asctime)s  %(levelname)-8s  %(name)s - %(message)s",
        },
    }

    return ExcelConfigBundle(
        settings=settings,
        site=site,
        profile=profile,
        warnings=warnings,
        excel_path=path,
    )


def _find_sheet(wb, keyword: str):
    for name in wb.sheetnames:
        if keyword.lower() in name.lower():
            return wb[name]
    raise KeyError(f"Sheet containing '{keyword}' not found. Available sheets: {wb.sheetnames}")


def _iter_kv(ws, start_row: int = 4) -> Iterable[Tuple[str, str]]:
    for row in ws.iter_rows(min_row=start_row, values_only=True):
        key = row[0]
        value = row[1] if len(row) > 1 else None
        key = _clean(key)
        if not key or key in HEADER_ROWS:
            continue
        yield key, _clean(value)


def _parse_uc4(ws) -> Tuple[Dict[str, str], Dict[str, str], Dict[str, str], Dict[str, str]]:
    section = ""
    pepguest_ids: Dict[str, str] = {}
    pepinet_ids: Dict[str, str] = {}
    pepguest_vars: Dict[str, str] = {}
    pepinet_vars: Dict[str, str] = {}

    for row in ws.iter_rows(min_row=4, values_only=True):
        key = _clean(row[0] if row else "")
        value = _clean(row[1] if row and len(row) > 1 else "")
        if not key:
            continue
        low = key.lower()
        if "pepguest" in low and "template ids" in low:
            section = "pepguest_ids"
            continue
        if "pepguest" in low and "input variables" in low:
            section = "pepguest_vars"
            continue
        if "pepinet" in low and "template ids" in low:
            section = "pepinet_ids"
            continue
        if "pepinet" in low and "input variables" in low:
            section = "pepinet_vars"
            continue
        if key in HEADER_ROWS:
            continue
        if section == "pepguest_ids":
            pepguest_ids[key] = value
        elif section == "pepinet_ids":
            pepinet_ids[key] = value
        elif section == "pepguest_vars":
            pepguest_vars[key] = value
        elif section == "pepinet_vars":
            pepinet_vars[key] = value

    return pepguest_ids, pepinet_ids, pepguest_vars, pepinet_vars


def _inject_uc3_logical_port(values: Dict[str, str], physical_port: str, vlan: int) -> Dict[str, str]:
    result = dict(values)
    if physical_port and vlan:
        logical_port = f"{physical_port}.{vlan}"
        for key in list(result):
            if "if-name" in key.lower() and not str(key).startswith("/0/"):
                result[key] = logical_port
    return result


def _drop_empty(data: Dict[str, str]) -> Dict[str, str]:
    return {k: v for k, v in data.items() if _clean(k) and _clean(v)}


def _first_non_empty(*values: str) -> str:
    for value in values:
        if value:
            return value
    return ""


def _to_int(value: str, default: int) -> int:
    try:
        return int(str(value).strip())
    except (TypeError, ValueError):
        return default


def _clean(value) -> str:
    if value is None:
        return ""
    return str(value).strip()
